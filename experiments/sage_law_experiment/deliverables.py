"""Writers and MiniMax-tolerant parsers for Harvey-style deliverable files."""

from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any


def strip_code_fences(text: str) -> str:
    return re.sub(r"```(?:json)?\s*\n?([\s\S]*?)\n?\s*```", r"\1", str(text or ""))


def strip_think_blocks(text: str) -> str:
    text = re.sub(r"<think>[\s\S]*?</think>", "", str(text or ""), flags=re.IGNORECASE)
    text = re.sub(r"<think>[\s\S]*", "", text, flags=re.IGNORECASE)
    return text.strip()


def extract_json_candidate(text: str) -> str | None:
    start = text.find("{")
    if start < 0:
        return None
    depth = 0
    in_string = False
    escape = False
    for index, ch in enumerate(text[start:], start=start):
        if escape:
            escape = False
            continue
        if ch == "\\" and in_string:
            escape = True
            continue
        if ch == '"':
            in_string = not in_string
            continue
        if in_string:
            continue
        if ch == "{":
            depth += 1
        elif ch == "}":
            depth -= 1
            if depth == 0:
                return text[start : index + 1]
    return text[start:]


def escape_newlines_inside_strings(raw: str) -> str:
    chars: list[str] = []
    in_string = False
    escape = False
    for ch in raw:
        if escape:
            chars.append(ch)
            escape = False
            continue
        if ch == "\\" and in_string:
            chars.append(ch)
            escape = True
            continue
        if ch == '"':
            chars.append(ch)
            in_string = not in_string
            continue
        if ch == "\n" and in_string:
            chars.append("\\n")
        else:
            chars.append(ch)
    return "".join(chars)


def repair_common_json(raw: str | None) -> str | None:
    if not raw:
        return None
    repaired = raw
    repaired = re.sub(r",\s*([}\]])", r"\1", repaired)
    repaired = repaired.replace("'", '"')
    repaired = re.sub(r"([{,]\s*)(\w+)\"\s*:", r'\1"\2":', repaired)
    repaired = re.sub(r"([{,]\s*)(\w+)\s*:", r'\1"\2":', repaired)
    return escape_newlines_inside_strings(repaired)


def repair_minimax_doubled_inner_quote_escapes(raw: str | None) -> str | None:
    """Repair MiniMax strings like ``\"Committee\\")`` inside JSON values.

    MiniMax occasionally emits an extra backslash before a quoted phrase's
    closing quote. In JSON, an even number of backslashes before ``"`` leaves
    the quote unescaped and prematurely closes the value string.
    """

    if not raw:
        return None
    return re.sub(r'\\{2,}"(?=\s*[^,}\]:])', r'\\"', raw)


def repair_invalid_backslash_escapes_inside_strings(raw: str | None) -> str | None:
    """Treat invalid JSON string escapes as literal backslashes.

    Legal deliverables often contain section references such as ``\\8.01``.
    MiniMax sometimes emits those as single JSON backslash escapes, which are
    invalid unless the following character is one of JSON's escape markers.
    """

    if not raw:
        return None
    valid_escapes = {'"', "\\", "/", "b", "f", "n", "r", "t", "u"}
    chars: list[str] = []
    in_string = False
    index = 0
    while index < len(raw):
        ch = raw[index]
        if ch == '"' and not in_string:
            in_string = True
            chars.append(ch)
            index += 1
            continue
        if ch == '"' and in_string:
            in_string = False
            chars.append(ch)
            index += 1
            continue
        if ch == "\\" and in_string:
            next_ch = raw[index + 1] if index + 1 < len(raw) else ""
            if next_ch in valid_escapes:
                chars.append(ch)
                if next_ch:
                    chars.append(next_ch)
                    index += 2
                else:
                    index += 1
                continue
            chars.append("\\\\")
            index += 1
            continue
        chars.append(ch)
        index += 1
    return "".join(chars)


def repair_truncated_json(text: str, drop_trailing_field: bool = False) -> str | None:
    match = re.search(r"\{[\s\S]*", str(text or ""))
    if not match:
        return None
    raw = match.group(0)
    if drop_trailing_field:
        raw = re.sub(r',\s*"[^"]*"?\s*:\s*"[^"]*$', "", raw)
        raw = re.sub(r',\s*"[^"]*$', "", raw)

    chars: list[str] = []
    open_braces = 0
    open_brackets = 0
    in_string = False
    escape = False
    for ch in raw:
        if escape:
            chars.append(ch)
            escape = False
            continue
        if ch == "\\" and in_string:
            chars.append(ch)
            escape = True
            continue
        if ch == '"':
            chars.append(ch)
            in_string = not in_string
            continue
        if not in_string:
            if ch == "{":
                open_braces += 1
            elif ch == "}":
                open_braces -= 1
            elif ch == "[":
                open_brackets += 1
            elif ch == "]":
                open_brackets -= 1
        chars.append("\\n" if ch == "\n" and in_string else ch)

    repaired = "".join(chars)
    if in_string:
        repaired += '"'
    repaired = re.sub(r",\s*$", "", repaired)
    repaired += "]" * max(0, open_brackets)
    repaired += "}" * max(0, open_braces)
    return repaired


def extract_json_object(text: str) -> dict[str, Any]:
    cleaned = strip_code_fences(strip_think_blocks(text)).strip()
    first_object = extract_json_candidate(cleaned)
    candidates = [
        cleaned,
        first_object,
        repair_minimax_doubled_inner_quote_escapes(first_object or cleaned),
        repair_invalid_backslash_escapes_inside_strings(first_object or cleaned),
        repair_invalid_backslash_escapes_inside_strings(
            repair_minimax_doubled_inner_quote_escapes(first_object or cleaned)
        ),
        repair_common_json(first_object or cleaned),
        repair_common_json(repair_minimax_doubled_inner_quote_escapes(first_object or cleaned)),
        repair_common_json(repair_invalid_backslash_escapes_inside_strings(first_object or cleaned)),
        repair_truncated_json(cleaned),
        repair_truncated_json(cleaned, drop_trailing_field=True),
    ]
    seen: set[str] = set()
    for candidate in candidates:
        if not candidate or candidate in seen:
            continue
        seen.add(candidate)
        try:
            parsed = json.loads(candidate)
        except json.JSONDecodeError:
            continue
        if isinstance(parsed, dict):
            return parsed
    raise ValueError("json_parse_failed")


def normalize_deliverable_source(parsed: dict[str, Any]) -> dict[str, Any]:
    source = parsed.get("deliverables", parsed)
    if isinstance(source, dict):
        return source
    if isinstance(source, list):
        normalized: dict[str, Any] = {}
        for item in source:
            if not isinstance(item, dict):
                continue
            filename = item.get("filename") or item.get("file") or item.get("name")
            content = item.get("content") or item.get("text") or item.get("body")
            if isinstance(filename, str) and content is not None:
                normalized[filename] = content
        return normalized
    if isinstance(source, str) and len(parsed) == 1:
        return {"content": source}
    return {}


def parse_deliverables(text: str, expected_filenames: list[str]) -> tuple[dict[str, str], str | None]:
    try:
        parsed = extract_json_object(text)
    except ValueError as exc:
        fallback = {filename: text for filename in expected_filenames}
        return fallback, f"deliverable_json_parse_failed: {exc}"

    source = normalize_deliverable_source(parsed)
    deliverables: dict[str, str] = {}
    missing = []
    for filename in expected_filenames:
        value = source.get(filename)
        if value is None:
            value = source.get(Path(filename).name)
        if value is None and len(expected_filenames) == 1:
            value = source.get("content") or source.get("text") or source.get("body")
        if value is None:
            missing.append(filename)
            deliverables[filename] = text
        elif isinstance(value, str):
            deliverables[filename] = value
        else:
            deliverables[filename] = json.dumps(value, indent=2, sort_keys=True)
    error = f"missing_deliverables: {', '.join(missing)}" if missing else None
    return deliverables, error


def write_deliverable(path: Path, content: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    content = xml_compatible_text(content)
    suffix = path.suffix.lower()
    if suffix == ".docx":
        write_docx(path, content)
    elif suffix == ".xlsx":
        write_xlsx(path, content)
    else:
        path.write_text(content + "\n")


def write_docx(path: Path, content: str) -> None:
    try:
        from docx import Document
    except Exception as exc:
        raise RuntimeError("python-docx is required to write .docx deliverables") from exc
    document = Document()
    for block in content.splitlines():
        line = block.strip()
        if not line:
            continue
        if line.startswith("# "):
            document.add_heading(line[2:].strip(), level=1)
        elif line.startswith("## "):
            document.add_heading(line[3:].strip(), level=2)
        elif line.startswith("### "):
            document.add_heading(line[4:].strip(), level=3)
        else:
            document.add_paragraph(line)
    document.save(path)


def write_xlsx(path: Path, content: str) -> None:
    try:
        from openpyxl import Workbook
    except Exception as exc:
        raise RuntimeError("openpyxl is required to write .xlsx deliverables") from exc
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Output"
    row_index = 1
    for line in content.splitlines():
        stripped = line.strip()
        if not stripped:
            row_index += 1
            continue
        if "|" in stripped:
            cells = [cell.strip() for cell in stripped.strip("|").split("|")]
        elif "," in stripped:
            cells = [cell.strip() for cell in stripped.split(",")]
        else:
            cells = [stripped]
        for col_index, cell in enumerate(cells, start=1):
            sheet.cell(row=row_index, column=col_index, value=cell)
        row_index += 1
    workbook.save(path)


def xml_compatible_text(content: str) -> str:
    """Remove control characters that Office XML writers cannot serialize."""

    return "".join(ch for ch in str(content) if ch in {"\t", "\n", "\r"} or ord(ch) >= 32)


def write_deliverable_set(
    output_dir: Path,
    deliverables: dict[str, str],
) -> None:
    for filename, content in deliverables.items():
        write_deliverable(output_dir / filename, content)


def combined_markdown(deliverables: dict[str, str]) -> str:
    parts = []
    for filename, content in deliverables.items():
        parts.append(f"# {filename}\n\n{content.strip()}")
    return "\n\n".join(parts).strip() + "\n"
